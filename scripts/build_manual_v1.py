from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Administrator\Documents\Deligo Triunfal, Lda")
OUT = ROOT / "docs" / "Manual_Corporativo_Operacional_Governacao_DeliGo_Central_v1.0.docx"
ASSET = ROOT / "docs" / "manual_assets"
ASSET.mkdir(parents=True, exist_ok=True)
IMG1 = Path(r"C:\Users\Administrator\Desktop\A TRABABA.png")
IMG2 = Path(r"C:\Users\Administrator\Desktop\A TRRRR.png")

BLUE = RGBColor(46, 116, 181); NAVY = RGBColor(31, 77, 120); DARK = RGBColor(28, 32, 39)
MUTED = RGBColor(91, 103, 117); LIGHT = "E8EEF5"; GREEN = "EAF7F0"; AMBER = "FFF4DB"; RED = "FDECEC"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn(f'w:{m}'))
        if node is None: node = OxmlElement(f'w:{m}'); mar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.first_child_found_in('w:tcW')
    if tcW is None: tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    if widths is None: widths = [9360//len(headers)]*len(headers)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_width(c,widths[i]); cell_margins(c); shade(c,LIGHT); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; r=p.add_run(h); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=NAVY
    trPr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trPr.append(hdr)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_width(cells[i],widths[i]); cell_margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(val)); r.font.size=Pt(9)
    return t

def para(doc, text='', bold=False, color=None, size=10.5, align=None, after=6):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.15
    if align is not None: p.alignment=align
    r=p.add_run(text); r.bold=bold; r.font.name='Calibri'; r.font.size=Pt(size)
    if color: r.font.color.rgb=color
    return p

def bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.1
    r=p.add_run(text); r.font.size=Pt(10.5); return p

def step(doc, n, title, detail):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.15); p.paragraph_format.first_line_indent=Inches(-.15); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(f"{n}. {title} — "); r.bold=True; r.font.color.rgb=NAVY; r.font.size=Pt(10.5)
    r=p.add_run(detail); r.font.size=Pt(10.5)

def callout(doc, label, text, fill=GREEN):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; c=t.cell(0,0); set_width(c,9360); cell_margins(c,140,180,140,180); shade(c,fill)
    p=c.paragraphs[0]; r=p.add_run(label.upper()+"\n"); r.bold=True; r.font.color.rgb=NAVY; r.font.size=Pt(9)
    r=p.add_run(text); r.font.size=Pt(10); return t

def heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.keep_with_next=True; return p

def flow_png(path, title, nodes, colors=None):
    w,h=1600,310; im=Image.new('RGB',(w,h),'white'); d=ImageDraw.Draw(im)
    try: font=ImageFont.truetype('arial.ttf',26); small=ImageFont.truetype('arial.ttf',20)
    except: font=ImageFont.load_default(); small=font
    d.text((40,20),title,fill='#1f4d78',font=font)
    x=40; y=115; boxw=240; gap=55
    for i,node in enumerate(nodes):
        col=(232,238,245) if not colors else colors[i%len(colors)]
        d.rounded_rectangle((x,y,x+boxw,y+90),radius=14,fill=col,outline='#2e74b5',width=3)
        tw=d.textbbox((0,0),node,font=small)[2]; d.text((x+(boxw-tw)//2,y+32),node,fill='#20242b',font=small)
        if i<len(nodes)-1:
            d.line((x+boxw+8,y+45,x+boxw+gap-12,y+45),fill='#2e74b5',width=4)
            d.polygon([(x+boxw+gap-12,y+45),(x+boxw+gap-26,y+36),(x+boxw+gap-26,y+54)],fill='#2e74b5')
        x+=boxw+gap
    im.save(path)

def add_image(doc, path, caption, width=6.4):
    doc.add_picture(str(path), width=Inches(width)); p=doc.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline = doc.inline_shapes[-1]._inline; docpr = inline.docPr; docpr.set('descr', caption); docpr.set('title', caption)
    c=para(doc,caption,align=WD_ALIGN_PARAGRAPH.CENTER,after=10); c.runs[0].italic=True; c.runs[0].font.size=Pt(8.5); c.runs[0].font.color.rgb=MUTED

def main():
    flow_png(ASSET/'flow_login.png','Acesso corporativo',['Telefone + senha','Validar backend','Sessão HQ','Visão Geral'])
    flow_png(ASSET/'flow_payroll.png','Payroll mensal',['Criar período','Elegíveis','Calcular','Rever','Aprovar','Pagar','Fechar'],['#e8eef5','#eaf7f0','#fff4db','#fff4db','#eaf7f0','#eaf7f0','#e8eef5'])
    flow_png(ASSET/'flow_kyc.png','KYC e Field',['Aquisição','Missão Field','Checklist','Decisão','Plataforma'],['#e8eef5','#fff4db','#fff4db','#eaf7f0','#e8eef5'])
    flow_png(ASSET/'flow_profit.png','Lucro e governação',['Receitas','Settlements','Custos','Lucro real','Reter / Reinvestir / Distribuir','Decisão 2/2'],['#e8eef5','#fff4db','#fff4db','#eaf7f0','#eaf7f0','#e8eef5'])

    doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    styles=doc.styles; normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
    for name,size,color,before,after in [('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,14,7),('Heading 3',12,NAVY,10,5)]:
        s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=color; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
    header=sec.header.paragraphs[0]; header.text='DElIGO TRIUNFAL · CENTRAL / HQ  |  MANUAL V1.0'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=MUTED
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=footer.add_run('Uso interno controlado · Versão 1.0 · 24 Agosto 2026'); rr.font.size=Pt(8); rr.font.color.rgb=MUTED

    para(doc,'DElIGO TRIUNFAL',bold=True,color=BLUE,size=12,after=3)
    para(doc,'MANUAL CORPORATIVO, OPERACIONAL E DE GOVERNAÇÃO',bold=True,color=DARK,size=23,after=4)
    para(doc,'DeliGo Central / HQ · Versão 1.0',color=NAVY,size=15,after=18)
    table(doc,['Controlo','Valor'],[['Proprietário','Direcção Executiva / Central HQ'],['Aplicação','Todas as plataformas e departamentos do ecossistema'],['Estado','Versão inicial controlada'],['Data','24 de Agosto de 2026'],['Classificação','Uso interno · informação operacional e corporativa']],widths=[2200,7160])
    para(doc,'Este manual consolida os princípios, fluxos, permissões e procedimentos definidos para a Central DeliGo Triunfal. Deve ser lido em conjunto com as políticas aprovadas, o registo de auditoria e a configuração efectiva do backend.',after=10)
    callout(doc,'Regra de controlo','A Central / HQ é o único produto administrativo interno. As plataformas operacionais utilizam a Conta Triunfal e enviam eventos, estados e movimentos para a Central; não se criam administrações paralelas.',LIGHT)
    doc.add_page_break()

    heading(doc,'Como utilizar este manual',1)
    para(doc,'O documento foi organizado como referência operacional: primeiro apresenta o modelo de governo e a arquitectura; depois descreve os módulos; por fim fornece procedimentos passo a passo por cargo, controlos de segurança e listas de verificação.')
    table(doc,['Leitor','Começar por','Responsabilidade central'],[
        ['CEO / SUPER_ADMIN','Secções 2, 3, 13 e 15','Governação, aprovações críticas, lucro e segurança'],
        ['RH','Secções 5 e 6','Pessoas, ciclo laboral, payroll e notificações'],
        ['Financeiro / Contabilidade','Secções 7, 8 e 9','Tesouraria, carteiras, diário, reconciliação e fecho'],
        ['Operações / Field','Secções 10 e 11','Missões, KYC, evidências, suporte e incidentes'],
        ['Tecnologia / Auditoria','Secções 3, 14 e 15','Disponibilidade, permissões, logs, incidentes e mudança']])
    heading(doc,'Índice operacional',2)
    for x in ['1. Finalidade e princípios','2. Arquitectura do ecossistema','3. Governo corporativo e decisões 2/2','4. Central, visão geral e navegação','5. Pessoas, cargos e permissões','6. RH, contratos e payroll','7. Finanças, wallets e tesouraria','8. Contabilidade, settlements e impostos','9. Lucro, reserve e distribuição 50/50','10. Field, KYC e localização','11. Suporte, aquisição e marketing','12. Compras, património, projectos e tarefas','13. BI, relatórios, notificações e auditoria','14. Perfil, tema, idioma e segurança','15. Procedimentos por cargo e checklists','Anexo A. Dicionário de estados','Anexo B. Evidência e controlo de versão']:
        bullet(doc,x)

    heading(doc,'1. Finalidade e princípios',1)
    para(doc,'A Central HQ fornece uma leitura única e governada da DeliGo Triunfal, Lda. O seu objectivo não é substituir os produtos operacionais, mas consolidar identidade, permissões, processos interdepartamentais, dados financeiros, riscos e decisões da empresa.')
    for title,detail in [('Uma Conta Triunfal','Uma pessoa mantém uma identidade central e pode ter múltiplos papéis por plataforma. Não duplicar contas administrativas.'),('Fonte de verdade','Estados financeiros, payroll, KYC e auditoria devem vir do backend autorizado. A interface nunca deve inventar valores.'),('Segregação de funções','Quem solicita não aprova a própria operação. Operações críticas e societárias exigem os dois CEOs.'),('Reversibilidade e auditoria','Períodos fechados não são editados directamente; correcções geram ajustes com histórico.'),('Tempo real com contexto','O stream actualiza a Central, mas cada cartão deve indicar a origem, instante de actualização e se o valor é real, histórico ou fictício de teste.')]:
        heading(doc,title,2); para(doc,detail)

    heading(doc,'2. Arquitectura do ecossistema',1)
    para(doc,'O ecossistema é composto por cinco plataformas operacionais ligadas a uma Central HQ: DeliGo, SARCA, DeliGo Market, Comunhão e Mestre24. O site institucional é um canal externo separado; não deve criar uma segunda administração interna.')
    table(doc,['Camada','Componentes','Regra de integração'],[
        ['Identidade','Conta Triunfal, telefone, sessão, papéis','Um utilizador; múltiplos vínculos; telefone como identificador de acesso'],
        ['Produtos','DeliGo, SARCA, Market, Comunhão, Mestre24','Cada produto conserva o seu fluxo operacional e publica eventos'],
        ['Central HQ','RH, Finanças, Contabilidade, Jurídico, Risco, IT e execução','Administração, aprovação, reconciliação e auditoria'],
        ['Dados','Supabase staging, RPCs, REST, ledger, auditoria','Service role apenas no servidor; cliente nunca recebe segredo'],
        ['Tempo real','Stream corporativo, actualizações sem cache, eventos','Cada alteração relevante deve reflectir-se na visão correspondente']])
    heading(doc,'Taxonomia dos produtos',2)
    table(doc,['Plataforma','Perfis externos principais','Exemplos de eventos para a Central'],[
        ['DeliGo','Cliente · Restaurante · Estafeta','Pedidos, comissões, entregas, KYC e settlements'],['SARCA','Rider · Driver','Viagens, motorista, passageiro, comissão e incidentes'],['DeliGo Market','Cliente · Vendedor · Empresa de logística','Compras, vendedores, entrega, split e carteira'],['Comunhão','Membro · Igreja','Membros, igrejas, eventos, donativos e subscrições'],['Mestre24','Cliente · Profissional','Pedidos, profissionais, subscrições e pagamentos']])

    heading(doc,'3. Governo corporativo e decisões 2/2',1)
    para(doc,'A DeliGo Triunfal tem dois CEOs com papel SUPER_ADMIN. A existência de dois administradores não elimina a segregação: decisões societárias, distribuição de lucros, operações críticas e gastos acima do limite exigem decisão de ambos.')
    table(doc,['Situação','Aprovação mínima','Nota de controlo'],[
        ['Até 2.500 MT','Responsável + Compras','Despesa pequena dentro do orçamento'],['2.501–10.000 MT','Responsável + Compras + Financeiro','Validação financeira obrigatória'],['10.001–50.000 MT','Financeiro + 1 CEO','Um CEO, sem auto-aprovação do beneficiário'],['Acima de 50.000 MT','Financeiro + 2 CEOs','Decisão 2/2'],['>20% da tesouraria operacional','Pelo menos 1 CEO','Regra de materialidade'],['>40% da tesouraria operacional','2 CEOs','Escala automática'],['Fora do orçamento','Subir um nível','Independentemente do valor'],['Operação societária/crítica','2 CEOs','Inclui lucro e operações de sócios']])
    callout(doc,'Independência','O CEO beneficiário de uma operação não pode ser a única pessoa a aprovar a própria operação. A decisão deve guardar autor, aprovadores, data, motivo e estado anterior/final.',AMBER)
    add_image(doc,ASSET/'flow_profit.png','Figura 1 — Fluxo de receitas, custos, lucro e decisão societária.')

    heading(doc,'4. Central, visão geral e navegação',1)
    para(doc,'A Visão Geral é o painel de entrada da Central. Deve apresentar cartões claros, cores funcionais, painel escuro de operações, mapa operacional e estado de integração. Os números devem ser calculados no backend e actualizados pelo stream corporativo.')
    add_image(doc,IMG1,'Figura 2 — Screenshot real fornecido da Visão Geral, com sidebar corporativa, cartões, estado operacional, painel escuro e mapa.')
    add_image(doc,IMG2,'Figura 3 — Screenshot real fornecido do Acesso Corporativo com identidade, telefone, palavra-passe, OTP e plataformas.')
    table(doc,['Área','Função','Resultado esperado'],[
        ['Visão geral','Resumo do grupo','Cartões e operações consolidadas; sem números históricos falsos'],['Plataformas','Estado das cinco plataformas','Ligação, sincronização, métricas e alertas'],['Mapa operacional','Eventos autorizados','Check-in/out e localização apenas quando permitido'],['Direcção Executiva','Decisões e governação','Deliberações, 2/2 e visão estratégica'],['Sistema','Homologação e integrações','Saúde técnica, versões, erros e sincronização']])
    step(doc,1,'Entrar','Abrir /login, usar telefone e palavra-passe; no primeiro acesso usar OTP e definir senha pessoal.')
    step(doc,2,'Validar','Aguardar a confirmação do backend. Em erro, não repetir operações financeiras; consultar notificações e auditoria.')
    step(doc,3,'Confirmar visão','Verificar que os cartões mostram origem e actualização, e que o perfil tem apenas os módulos autorizados.')
    step(doc,4,'Navegar','Usar o menu lateral por domínio; a rota e o estado devem permanecer no browser após actualização.')

    heading(doc,'5. Pessoas, cargos e permissões',1)
    para(doc,'Os trabalhadores internos não são os mesmos que os perfis externos das plataformas. A Central utiliza uma estrutura hierárquica de três níveis: Direcção Executiva, responsáveis de departamento e colaboradores/especialistas.')
    table(doc,['Nível','Exemplo','Poderes'],[['Direcção Executiva','2 CEOs · SUPER_ADMIN','Visão total, decisões críticas, governo e segurança'],['Responsável','Director RH, Financeiro, IT','Gerir o domínio; aprovar dentro do âmbito'],['Colaborador','Assistente RH, analista, agente Field','Executar tarefas e acções autorizadas; sem auto-aprovação']])
    heading(doc,'Departamentos internos',2)
    para(doc,'A Central organiza pessoas e processos em: Direcção Executiva, Pessoas/RH, Finanças, Contabilidade, Jurídico, Conformidade e Risco, Tecnologia/IT, Compras, Património, Marketing, Operações de Campo, Apoio ao Cliente, Aquisição de Parceiros, Projectos e Tarefas, Apoio Interno, Notificações, BI/Análises, Relatórios, Auditoria e Homologação.')
    step(doc,1,'Criar trabalhador','Pessoas → Novo trabalhador; registar nome, telefone +258, cargo, departamento e role. Deixar activação desmarcada quando o cadastro exigir aprovação.')
    step(doc,2,'Aprovar acesso','O Super Admin verifica identidade, departamento e necessidade de acesso; aprova ou mantém convidado.')
    step(doc,3,'Atribuir permissões','Conceder apenas códigos necessários. FIELD_KYC_REVIEW e FIELD_KYC_DECIDE são permissões diferentes.')
    step(doc,4,'Auditar','Toda alteração de estado, supervisor ou override de permissão deve aparecer no histórico de acessos e auditoria.')

    heading(doc,'6. RH, contratos e payroll',1)
    para(doc,'O RH é o proprietário do ciclo laboral: cadastro, admissão, contrato, documentos, presenças, férias, payroll, avaliação, formação, promoção, transferência e desligamento.')
    add_image(doc,ASSET/'flow_payroll.png','Figura 4 — Fluxo normativo do payroll mensal.')
    heading(doc,'Payroll por período',2)
    para(doc,'A folha salarial é consequência de um período salarial. Primeiro cria-se o período; depois o sistema identifica trabalhadores elegíveis pela data efectiva, contrato e estado laboral.')
    for x in ['Agosto 2026 · 01/08/2026–31/08/2026 · Rascunho','Identificar trabalhadores activos e elegíveis no período','Carregar salário base do contrato oficial','Registar dias trabalhados, faltas, horas extra, bónus, subsídios, adiantamentos e descontos','Calcular bruto, descontos e líquido','Rever RH; aprovar Direcção/Financeiro; financiar Payroll Wallet; pagar; gerar comprovativos; fechar','Após fecho, correcção apenas por ajuste e auditoria']:
        bullet(doc,x)
    table(doc,['Campo','Regra'],[['Salário base','Vem do contrato/perfil salarial oficial; não é redigitado todos os meses'],['Dias trabalhados','Usar datas e presenças elegíveis'],['Faltas/ausências','Registar motivo, período e aprovação quando aplicável'],['Horas extra','Aplicar política e evidência; nunca inferir automaticamente'],['Descontos','Apenas autorizados e auditáveis'],['Estado','Rascunho → Calculado → Em revisão → Aprovado → Financiado → Pago → Fechado']])
    callout(doc,'Protecção','Um mês fraco não reduz automaticamente salários. A empresa fecha o período, calcula o lucro real e decide retenção, reinvestimento ou distribuição separadamente.',RED)

    heading(doc,'7. Finanças, wallets e tesouraria',1)
    para(doc,'Finanças controla dinheiro: contas, transacções, despesas, receitas, reservas, pedidos de pagamento, payroll wallet e políticas de aprovação. A Contabilidade regista e explica os efeitos contabilísticos.')
    heading(doc,'Modelo correcto de wallets de teste',2)
    para(doc,'Cada utilizador tem uma carteira central única. As cinco plataformas usam o mesmo wallet_id e guardam platform_slug no movimento. Para os três utilizadores de teste, o modelo esperado é: 3 carteiras centrais × 20.000 MZN = 60.000 MZN. Não duplicar para 15 carteiras.')
    table(doc,['Conta de teste','Saldo fictício','Regra'],[['841130015','20.000 MZN','Uma carteira central; vínculos por plataforma'],['827688695','20.000 MZN','Uma carteira central; vínculos por plataforma'],['877688695','20.000 MZN','Uma carteira central; vínculos por plataforma'],['Total','60.000 MZN','fictitious=true e test_mode=true; não é lucro real']])
    step(doc,1,'Registar movimento','Criar evento com wallet_id, plataforma de origem, tipo, valor, moeda, referência e idempotency key.')
    step(doc,2,'Validar saldo','Verificar ledger e saldo disponível; nunca confiar apenas no valor apresentado no cliente.')
    step(doc,3,'Reconciliar','Confirmar evidência de banco, M-Pesa ou e-Mola antes de marcar como publicado/pago.')
    step(doc,4,'Auditar','Guardar autor, fonte, timestamps, estado, reversão e relação com settlement.')

    heading(doc,'8. Contabilidade, settlements e impostos',1)
    para(doc,'Finanças paga e recebe; Contabilidade classifica, reconcilia, fecha períodos, calcula obrigações e produz demonstrações. O Plano de Contas deve existir antes de depender dos dashboards.')
    table(doc,['Grupo','Exemplos'],[['Activo','Caixa e Bancos; Contas a Receber'],['Passivo','Contas a Pagar; obrigação com parceiros; impostos'],['Resultado','Receitas por plataforma; despesas operacionais; payroll'],['Capital','Capital Social; Resultados Acumulados; Profit Reserve']])
    para(doc,'Settlement não é receita. Se entrarem 1.000 MT, sendo 900 MT do parceiro e 100 MT de comissão DeliGo, o lançamento conceptual é: Caixa/Banco +1.000; obrigação com parceiro +900; receita DeliGo +100.')
    step(doc,1,'Abrir período','Contabilidade cria Agosto 2026: Aberto.')
    step(doc,2,'Publicar movimentos','Pagamentos confirmados e movimentos reconciliados geram diário e razão de forma idempotente.')
    step(doc,3,'Reconciliar','Preparador associa transacção à evidência; segundo utilizador valida.')
    step(doc,4,'Fechar','Período passa a Em fechamento e depois Fechado; ajustes futuros mantêm o histórico.')

    heading(doc,'9. Lucro, Profit Reserve e distribuição 50/50',1)
    para(doc,'O lucro é um resultado calculado, não um saldo bancário nem o volume bruto das plataformas. A Central deve separar receitas, settlements, custos, payroll, impostos e despesas antes de apurar o resultado líquido.')
    table(doc,['Fase','Tratamento'],[['Receitas','Comissões e receitas próprias reconhecidas'],['Settlements','Obrigações devidas a parceiros; não são lucro'],['Custos','Payroll, fornecedores, operação e impostos'],['Lucro real','Receita própria menos custos e obrigações reconhecidas'],['Destino','Reter, reinvestir ou enviar para Profit Reserve'],['Distribuição','Só depois de período fechado e decisão 2/2'],['Sócios','Distribuição aprovada em 50% / 50%; não automática para wallets']])
    add_image(doc,ASSET/'flow_profit.png','Figura 5 — Governação do lucro e distribuição.')
    callout(doc,'Não automático','O sistema nunca deve reduzir salários por mês fraco nem enviar lucro automaticamente para wallets pessoais. A decisão deve ser explícita, aprovada e auditada.',AMBER)

    heading(doc,'10. Field, KYC e localização',1)
    para(doc,'Operações de Campo trabalha com Missões. Checklist, evidências, formulário, incidentes e relatório são componentes da missão, não entidades paralelas.')
    add_image(doc,ASSET/'flow_kyc.png','Figura 6 — Fluxo de aquisição, Field, KYC e decisão.')
    table(doc,['Estado da missão','Evidência'],[['Criada','Objectivo, plataforma, alvo, agente, data e prioridade'],['Atribuída / Aceite','Agente e aceitação registados'],['Check-in','Hora e coordenada autorizada'],['Em execução','Checklist, observações e evidências'],['Check-out','Hora e coordenada autorizada'],['Relatório / Concluída','Resultado, incidentes e decisão']])
    heading(doc,'KYC',2)
    para(doc,'O dashboard deve distinguir Recebidos, Pendentes, Em análise, Correção solicitada, Aprovados, Rejeitados e Escalados. FIELD_KYC_REVIEW permite analisar; FIELD_KYC_DECIDE permite aprovar/rejeitar.')
    step(doc,1,'Abrir caso','Field recebe KYC da plataforma e valida identidade, documentos e checklist.')
    step(doc,2,'Decidir','APROVAR, PEDIR CORRECÇÃO ou REJEITAR. Motivo obrigatório para correcção/rejeição.')
    step(doc,3,'Escalar','Documento adulterado, divergência grave ou suspeita de fraude → Conformidade e Risco.')
    step(doc,4,'Publicar','Só após decisão válida o resultado volta à plataforma de origem.')
    para(doc,'Localização contínua não é padrão: usar check-in e check-out autorizados; rastreamento permanente somente quando necessário, permitido e informado.')

    heading(doc,'11. Suporte, aquisição e marketing',1)
    para(doc,'Apoio ao Cliente recebe ocorrências e acompanha SLA; Aquisição gere leads, parceiros, visitas e onboarding; Marketing gere campanhas, conteúdo e activaçōes. Os módulos comunicam por processos e notificações, sem saltar controlos.')
    table(doc,['Origem','Destino','Quando'],[['Plataformas','Field','KYC, verificação e visita'],['Aquisição','Field','Onboarding e visita presencial'],['Marketing','Field','Campanha ou activação no terreno'],['Field','Conformidade','Caso suspeito/excepcional'],['Field','Finanças','Despesa/reembolso da missão'],['Field','Suporte','Ocorrência que exige acompanhamento'],['Field','Plataforma','Decisão final de KYC']])
    step(doc,1,'Registar caso','Criar ticket/lead/missão com plataforma, responsável, prioridade e prazo.')
    step(doc,2,'Atribuir','Definir owner e equipa; não deixar casos críticos sem responsável.')
    step(doc,3,'Executar','Registar mensagens, visitas, evidências e decisões no histórico.')
    step(doc,4,'Encerrar','Confirmar resultado com o requerente; fechar apenas com resolução ou motivo documentado.')

    heading(doc,'12. Compras, património, projectos e tarefas',1)
    para(doc,'Compras segue Pedido → Aprovação → Fornecedor → Recepção → Pagamento → Contabilidade. Património controla activo, custódia, manutenção e baixa. Projectos e Tarefas registam execução, dependências, responsáveis e marcos.')
    table(doc,['Fluxo','Estados mínimos'],[['Compra','Rascunho → Submetido → Em aprovação → Aprovado → Pedido → Recebido → Pago'],['Activo','Registado → Em stock → Atribuído → Manutenção → Reformado'],['Projecto','Planeado → Activo → Em risco/Bloqueado → Concluído'],['Tarefa','Planeada → Atribuída → Em curso → Em revisão → Concluída']])
    step(doc,1,'Planear','Definir objectivo, departamento, dono, prioridade, orçamento e prazo.')
    step(doc,2,'Aprovar','Aplicar limite financeiro, materialidade e orçamento.')
    step(doc,3,'Executar','Atribuir, acompanhar progresso, anexar evidências e registar mudanças.')
    step(doc,4,'Controlar','BI, Relatórios e Auditoria devem conseguir explicar quem fez o quê e quando.')

    heading(doc,'13. BI, relatórios, notificações e auditoria',1)
    para(doc,'BI / Análises é o painel de inteligência: indicadores, tendências, filtros por plataforma/departamento e alertas. Relatórios são artefactos formais; Auditoria é o registo append-only da actividade.')
    table(doc,['Módulo','Regra'],[['BI','Explicar origem, período, filtro e fórmula; separar teste de produção'],['Relatórios','Snapshots identificáveis, versão, proprietário e período'],['Notificações','Categorias, prioridade, estado de leitura e entrega; não duplicar'],['Auditoria','Autor, acção, módulo, alvo, resultado, antes/depois, IP/device quando disponível'],['Aprovações','Fila filtrada por permissão; decisão e comentário obrigatórios quando requerido']])
    callout(doc,'Erro comum','Unexpected end of JSON input indica normalmente resposta vazia ou truncada num serviço comum. O cliente deve validar status/content-type antes de fazer JSON.parse e apresentar erro accionável.',RED)
    step(doc,1,'Detectar','Abrir Registo de Auditoria ou Tecnologia/IT e identificar request, rota e timestamp.')
    step(doc,2,'Preservar','Não apagar o log nem repetir a transacção sem idempotency key.')
    step(doc,3,'Corrigir','Ajustar serviço partilhado/API e adicionar teste de resposta vazia, erro e sucesso.')
    step(doc,4,'Confirmar','Reexecutar fluxo ponta a ponta e guardar evidência da correcção.')

    heading(doc,'14. Perfil, tema, idioma e segurança',1)
    para(doc,'Meu Perfil contém Perfil, Contactos, Pagamento, Segurança, Preferências, Notificações, Documentos e Acessos. Preferências individuais não precisam de aprovação de RH ou CEOs.')
    table(doc,['Preferência','V1'],[['Aparência','Sistema · Claro · Escuro'],['Idioma','Português (Moçambique) · English'],['Formato de data','Automático por idioma/região'],['Fuso horário','Africa/Maputo'],['Persistência','Guardar por utilizador no backend; tema rápido no topo']])
    para(doc,'Tema não altera semântica funcional: Aprovado, Pendente, Rejeitado, Crítico e Em risco devem continuar distinguíveis por cor, texto e acessibilidade. i18n deve cobrir menus, botões, erros, estados e notificações; nomes próprios e documentos não são traduzidos.')
    table(doc,['Controlo','Procedimento'],[['Sessão','Cookie seguro, expiração, logout e verificação de actor'],['Segredos','Service role somente no servidor; nunca em cliente ou screenshot'],['Permissões','Deny by default; escopos por módulo e plataforma'],['Dados pessoais','RH e salários restritos; documentos com acesso auditado'],['Mudanças','Migração versionada, revisão, teste e plano de rollback'],['Incidentes','Registar, conter, comunicar, corrigir e fazer post-mortem']])

    heading(doc,'15. Procedimentos por cargo e checklists',1)
    roles=[
        ('CEO / SUPER_ADMIN',['Rever dashboard e saúde das integrações','Aprovar matérias 2/2 e distribuição de lucro','Validar alterações críticas, acessos e incidentes','Rever auditoria e decisões pendentes']),
        ('Responsável de RH',['Criar/actualizar trabalhador e ciclo laboral','Validar contrato e elegibilidade de payroll','Calcular e rever folha do período','Encaminhar aprovação e notificação']),
        ('Financeiro',['Validar tesouraria e orçamento','Rever compras, pagamentos e payroll wallet','Confirmar reconciliação e evidência','Preparar lucro real sem confundir settlements']),
        ('Contabilidade',['Manter Plano de Contas','Publicar diário e razão idempotentes','Reconciliar e fechar período','Produzir balancete, DRE, balanço e fluxo de caixa']),
        ('Field / KYC',['Aceitar missão e fazer check-in','Executar checklist e recolher evidência','Analisar/decidir apenas com permissão','Escalar fraude ou divergência']),
        ('Apoio ao Cliente',['Registar ticket e prioridade','Atribuir owner e SLA','Responder e anexar evidência','Fechar com resolução ou escalamento']),
        ('Tecnologia / IT',['Monitorizar runtime, APIs e stream','Investigar erros e JSON vazio','Controlar deploy/migração e rollback','Preservar logs e segurança']),
        ('Auditoria / Risco',['Rever eventos críticos','Testar segregação de funções','Verificar ajustes após fecho','Emitir achados e acompanhar correcção'])]
    for role,items in roles:
        heading(doc,role,2)
        for i,x in enumerate(items,1): bullet(doc,f'{i}. {x}')
    heading(doc,'Checklist de aceitação V1.0',2)
    for x in ['Acesso corporativo funciona com telefone, palavra-passe e OTP','Visão Geral conserva sidebar, cores, cartões, painel escuro e mapa','Cartões distinguem produção, histórico e dados fictícios de teste','Os três utilizadores de teste têm uma carteira central cada','Payroll só é gerado a partir de período e contrato oficial','Período fechado não aceita edição directa','Settlement não entra como receita da DeliGo','Lucro só pode ser distribuído após fecho e decisão 2/2','Field consegue executar missão e KYC dentro da permissão','Compras aplica limites 2.500 / 10.000 / 50.000 MT e materialidade','Notificações, tarefas, aprovações e auditoria conservam histórico','Tema/idioma persistem por utilizador','Service role e segredos não aparecem no cliente','Todos os fluxos críticos têm evidência PASS/FAIL/BLOCKED']:
        bullet(doc,'☐ '+x)

    heading(doc,'Anexo A. Dicionário de estados',1)
    table(doc,['Estado','Significado'],[['Rascunho','Pode ser completado pelo proprietário'],['Submetido','A aguardar validação/aprovação'],['Em revisão','Há análise em curso'],['Aprovado','Decisão positiva com actor e timestamp'],['Pendente','Falta informação, documento ou aprovação'],['Em risco','Há condição que pode impedir prazo/resultado'],['Bloqueado','Não pode avançar sem intervenção'],['Fechado','Período/processo concluído; alterações por ajuste'],['Rejeitado','Decisão negativa com motivo obrigatório'],['Escalado','Enviado a autoridade/departamento superior']])
    heading(doc,'Anexo B. Evidência e controlo de versão',1)
    para(doc,'Screenshots: as Figuras 2 e 3 são imagens reais fornecidas da Central. Os fluxogramas são representações normativas V1.0 dos procedimentos descritos. A implementação deve ser confirmada no ambiente staging com testes autenticados, respostas HTTP, permissões e dados reais de teste; uma tela ou build isolado não prova o fluxo ponta a ponta.')
    table(doc,['Versão','Data','Alteração','Responsável'],[['1.0','24/08/2026','Manual inicial consolidado: arquitectura, operações, governo, finanças, RH, segurança e procedimentos','DeliGo Central / HQ']])
    callout(doc,'Aprovação','Este manual é uma base operacional V1.0. Políticas legais, fiscais, laborais e societárias devem ser validadas pela Direcção e pelos responsáveis competentes antes de produção.',AMBER)
    doc.core_properties.title='Manual Corporativo, Operacional e de Governação - DeliGo Central HQ V1.0'; doc.core_properties.subject='Manual interno'; doc.core_properties.author='DeliGo Triunfal Central'; doc.core_properties.comments='Versão 1.0 controlada'
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)

if __name__=='__main__': main()
